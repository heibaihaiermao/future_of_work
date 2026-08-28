from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles import Alignment


HEADERS = [
    "Skill Gap Name",
    "Explanation",
    "Justification",
]


def export_glossary(entries, output_file):

    wb = Workbook()

    ws = wb.active
    ws.title = "Glossary"

    for col, header in enumerate(HEADERS, start=1):

        cell = ws.cell(
            row=1,
            column=col,
            value=header
        )

        cell.font = Font(bold=True)

    row = 2

    for entry in entries:

        ws.cell(row=row, column=1, value=entry["skill_gap"])
        ws.cell(row=row, column=2, value=entry["explanation"])
        ws.cell(row=row, column=3, value=entry["justification"])

        row += 1

    widths = {
        "A": 50,
        "B": 100,
        "C": 100,
    }

    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    for row_cells in ws.iter_rows():

        for cell in row_cells:

            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top"
            )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{ws.max_row}"

    wb.save(output_file)